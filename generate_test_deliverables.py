"""Batch Test Execution & Deliverable Generator for SEO Audit Agent.

Runs test suites for Q1 (SEO Audit), Q2 (NAP Check), and Q3 (Grounded QA) across target websites.
Generates:
1. Individual & combined JSON outputs in `outputs/`
2. Inputs summary documentation in Excel (.xlsx), Word (.docx), and Text (.txt) formats
3. A ZIP bundle of all generated outputs.
"""

from __future__ import annotations

import json
import os
import shutil
import sys
import zipfile
from dataclasses import asdict, dataclass
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import docx
from docx.shared import Inches, Pt, RGBColor

from app.main import build_llm_generate, run_pipeline, write_outputs


@dataclass
class TestCase:
    id: str
    target_url: str
    question_type: str  # "Q1 (SEO Findings)", "Q2 (NAP Consistency)", "Q3 (Grounded QA)", "All (Q1+Q2+Q3)"
    question: str | None
    max_pages: int = 10
    max_depth: int = 2
    llm_provider: str = "nvidia"
    description: str = ""


TEST_CASES: list[TestCase] = [
    TestCase(
        id="TC01",
        target_url="https://webflow.com/",
        question_type="Q3 (Grounded QA)",
        question="Can webflow is used for website development?",
        description="Verify webflow website development capabilities in QA agent.",
    ),
    TestCase(
        id="TC02",
        target_url="https://www.python.org",
        question_type="Q3 (Grounded QA)",
        question="What is python?",
        description="Verify general definition/overview answering for python.org.",
    ),
    TestCase(
        id="TC03",
        target_url="https://ccbp.in/",
        question_type="Q2 (NAP Consistency)",
        question=None,
        description="Check Name, Address, Phone (NAP) consistency across ccbp.in pages.",
    ),
    TestCase(
        id="TC04",
        target_url="https://www.niatindia.com/",
        question_type="Q2 (NAP Consistency)",
        question=None,
        description="Check Name, Address, Phone (NAP) consistency across niatindia.com pages.",
    ),
    TestCase(
        id="TC05",
        target_url="https://www.thehudsonkitchen.com/",
        question_type="Q2 (NAP Consistency)",
        question=None,
        description="Check Name, Address, Phone (NAP) consistency across local business pages of Hudson Kitchen.",
    ),
    TestCase(
        id="TC06",
        target_url="https://ccbp.in/",
        question_type="Q3 (Grounded QA)",
        question="What is this domain for?",
        description="Extract evidence and explain core purpose of ccbp.in domain.",
    ),
    TestCase(
        id="TC07",
        target_url="https://www.niatindia.com/",
        question_type="Q3 (Grounded QA)",
        question="What programs or courses are offered by NIAT?",
        description="Extract course/program offerings with grounded citations from NIAT India.",
    ),
    TestCase(
        id="TC08",
        target_url="https://www.thehudsonkitchen.com/",
        question_type="Q3 (Grounded QA)",
        question="Where is Hudson Kitchen located and what services do they provide?",
        description="Answer location and service details for Hudson Kitchen using extracted page text.",
    ),
    TestCase(
        id="TC09",
        target_url="https://webflow.com/",
        question_type="Q1 (SEO Audit Findings)",
        question=None,
        description="Run technical & on-page SEO audit findings against webflow.com homepage & subpages.",
    ),
    TestCase(
        id="TC10",
        target_url="https://www.python.org",
        question_type="Q1 (SEO Audit Findings)",
        question=None,
        description="Run technical & on-page SEO audit findings against python.org.",
    ),
]


def run_batch_tests(base_output_dir: str = "outputs") -> dict[str, Any]:
    os.makedirs(base_output_dir, exist_ok=True)
    batch_dir = os.path.join(base_output_dir, "batch_runs")
    os.makedirs(batch_dir, exist_ok=True)

    llm_key = os.environ.get("LLM_API_KEY")
    all_results_summary = []

    print("=" * 70)
    print(f"Starting execution of {len(TEST_CASES)} test cases...")
    print("=" * 70)

    for i, tc in enumerate(TEST_CASES, 1):
        print(f"\n[{i}/{len(TEST_CASES)}] Running {tc.id} ({tc.question_type}) on {tc.target_url}")
        if tc.question:
            print(f"  Question: {tc.question!r}")

        test_out_dir = os.path.join(batch_dir, f"{tc.id}_{tc.target_url.replace('https://', '').replace('/', '_')}")
        os.makedirs(test_out_dir, exist_ok=True)

        llm_gen = build_llm_generate(tc.llm_provider, llm_key)

        try:
            res = run_pipeline(
                url=tc.target_url,
                question=tc.question,
                max_pages=tc.max_pages,
                max_depth=tc.max_depth,
                timeout=12,
                llm_generate=llm_gen,
            )
            written_paths = write_outputs(res, test_out_dir)

            with open(os.path.join(test_out_dir, "audit.json"), "r", encoding="utf-8") as f:
                audit_data = json.load(f)
            with open(os.path.join(test_out_dir, "nap_report.json"), "r", encoding="utf-8") as f:
                nap_data = json.load(f)

            ans_data = None
            ans_path = os.path.join(test_out_dir, "answer.json")
            if os.path.exists(ans_path):
                with open(ans_path, "r", encoding="utf-8") as f:
                    ans_data = json.load(f)

            summary_item = {
                "test_id": tc.id,
                "target_url": tc.target_url,
                "question_type": tc.question_type,
                "question": tc.question,
                "max_pages": tc.max_pages,
                "max_depth": tc.max_depth,
                "pages_crawled": len(res.pages),
                "pages_fetched": res.pages_fetched,
                "findings_count": len(audit_data),
                "nap_entries_count": len(nap_data),
                "qa_answered": ans_data is not None and ans_data.get("url") is not None,
                "output_dir": test_out_dir,
                "audit": audit_data,
                "nap_report": nap_data,
                "answer": ans_data,
            }
            all_results_summary.append(summary_item)

            print(f"  ✓ Crawled {len(res.pages)} pages. Q1 Findings: {len(audit_data)}, Q2 NAP: {len(nap_data)}")
            if ans_data:
                status_str = f"Found on {ans_data.get('url')}" if ans_data.get("url") else "No grounded answer"
                print(f"  ✓ Q3 Answer: {status_str}")

        except Exception as exc:
            print(f"  ❌ Error running test case {tc.id}: {exc}")

    # Save combined outputs JSON
    combined_json_path = os.path.join(base_output_dir, "all_test_outputs_combined.json")
    with open(combined_json_path, "w", encoding="utf-8") as f:
        json.dump(all_results_summary, f, indent=2, ensure_ascii=False)
    print(f"\nSaved combined outputs JSON: {combined_json_path}")

    return {"test_cases": TEST_CASES, "results": all_results_summary}


def generate_excel_inputs(test_cases: list[TestCase], file_path: str) -> None:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styling definitions
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    title_font = Font(name="Calibri", size=13, bold=True, color="1F4E78")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    categories = [
        ("Q1 - SEO Audit Findings", "Q1", "Question 1: Technical & On-Page SEO Audit Findings Inputs"),
        ("Q2 - NAP Audit", "Q2", "Question 2: Name, Address, Phone (NAP) Consistency Audit Inputs"),
        ("Q3 - Grounded QA", "Q3", "Question 3: Evidence-Grounded Question Answering (QA) Inputs"),
        ("All Inputs Overview", "ALL", "Complete Matrix of All Tested Inputs Across Questions"),
    ]

    headers = [
        "Test ID",
        "Target Website URL",
        "User Question (Q3)",
        "Max Pages",
        "Max Depth",
        "LLM Provider",
        "Description & Objective",
    ]

    for sheet_title, filter_code, title_desc in categories:
        ws = wb.create_sheet(title=sheet_title)
        ws.append([title_desc])
        ws.cell(1, 1).font = title_font
        ws.append([])  # Spacer

        ws.append(headers)
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=3, column=col_num)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

        filtered_tc = (
            test_cases
            if filter_code == "ALL"
            else [tc for tc in test_cases if filter_code in tc.question_type]
        )

        for tc in filtered_tc:
            row = [
                tc.id,
                tc.target_url,
                tc.question or "N/A (Audit / NAP Only)",
                tc.max_pages,
                tc.max_depth,
                tc.llm_provider.upper(),
                tc.description,
            ]
            ws.append(row)
            curr_row = ws.max_row
            for col_num in range(1, len(headers) + 1):
                c = ws.cell(row=curr_row, column=col_num)
                c.font = data_font
                c.border = thin_border
                if col_num in (1, 4, 5, 6):
                    c.alignment = Alignment(horizontal="center", vertical="center")

        widths = [10, 34, 45, 12, 12, 14, 55]
        for i, col_letter in enumerate(["A", "B", "C", "D", "E", "F", "G"]):
            ws.column_dimensions[col_letter].width = widths[i]

    wb.save(file_path)
    print(f"Generated Excel inputs file (with Q1, Q2, Q3 tabs): {file_path}")


def generate_docx_inputs(test_cases: list[TestCase], file_path: str) -> None:
    doc = docx.Document()
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    # Title Header
    p_title = doc.add_paragraph()
    run_title = p_title.add_run("Evidence-Grounded SEO Audit Agent")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(31, 78, 120)

    p_sub = doc.add_paragraph()
    run_sub = p_sub.add_run("Test Inputs Matrix (Separated by Question 1, Question 2, and Question 3)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph("This document details all test input parameters evaluated, strictly separated by Question category.")

    sections = [
        ("Question 1 (Q1) — Technical & On-Page SEO Audit Findings", "Q1"),
        ("Question 2 (Q2) — Local SEO & NAP Consistency Audit", "Q2"),
        ("Question 3 (Q3) — Evidence-Grounded Question Answering (QA)", "Q3"),
    ]

    for title, code in sections:
        h = doc.add_heading(title, level=2)
        for run in h.runs:
            run.font.color.rgb = RGBColor(31, 78, 120)

        filtered = [tc for tc in test_cases if code in tc.question_type]
        if not filtered:
            continue

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"

        hdr_cells = table.rows[0].cells
        hdr_titles = ["ID", "Target URL", "Question (Q3)", "Crawl Limits", "Objective & Description"]
        for i, t in enumerate(hdr_titles):
            hdr_cells[i].text = t
            shd = parse_xml(r'<w:shd {} w:fill="1F4E78"/>'.format(nsdecls('w')))
            hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
            for p in hdr_cells[i].paragraphs:
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = RGBColor(255, 255, 255)

        for tc in filtered:
            row_cells = table.add_row().cells
            row_cells[0].text = tc.id
            row_cells[1].text = tc.target_url
            row_cells[2].text = tc.question or "N/A"
            row_cells[3].text = f"Pages: {tc.max_pages}\nDepth: {tc.max_depth}"
            row_cells[4].text = tc.description

            for cell in row_cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.name = "Arial"
                        run.font.size = Pt(8.5)

        doc.add_paragraph("")  # Spacer between sections

    doc.save(file_path)
    print(f"Generated Word inputs file (separated by Q1, Q2, Q3): {file_path}")


def generate_txt_inputs(test_cases: list[TestCase], file_path: str) -> None:
    lines = [
        "=========================================================================================",
        "            EVIDENCE-GROUNDED SEO AUDIT AGENT — TEST INPUTS (SEPARATED BY QUESTION)",
        "=========================================================================================",
        "Summary of input parameters tested independently, separated into Question 1, Question 2,",
        "and Question 3.",
        "=========================================================================================\n",
    ]

    sections = [
        ("SECTION 1: QUESTION 1 (Q1) — TECHNICAL & ON-PAGE SEO AUDIT FINDINGS", "Q1"),
        ("SECTION 2: QUESTION 2 (Q2) — LOCAL SEO & NAP CONSISTENCY AUDIT", "Q2"),
        ("SECTION 3: QUESTION 3 (Q3) — EVIDENCE-GROUNDED QUESTION ANSWERING (QA)", "Q3"),
    ]

    for sec_title, code in sections:
        lines.append("=" * 89)
        lines.append(f" {sec_title}")
        lines.append("=" * 89 + "\n")

        filtered = [tc for tc in test_cases if code in tc.question_type]
        for tc in filtered:
            lines.append(f"  Test ID       : {tc.id}")
            lines.append(f"  Target URL    : {tc.target_url}")
            lines.append(f"  Question (Q3) : {tc.question or 'N/A (Audit / NAP check only)'}")
            lines.append(f"  Crawl Limits  : Max Pages = {tc.max_pages}, Max Depth = {tc.max_depth}, Timeout = 12s")
            lines.append(f"  LLM Provider  : {tc.llm_provider.upper()}")
            lines.append(f"  Description   : {tc.description}")
            lines.append("  " + "-" * 85 + "\n")
        lines.append("\n")

    with open(file_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    print(f"Generated Text inputs file (separated by Q1, Q2, Q3): {file_path}")


def create_outputs_zip(base_output_dir: str = "outputs", zip_filename: str = "outputs/all_test_outputs.zip") -> None:
    abs_zip_path = os.path.abspath(zip_filename)
    print(f"\nCreating ZIP archive: {abs_zip_path}")

    with zipfile.ZipFile(abs_zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for root, dirs, files in os.walk(base_output_dir):
            for file in files:
                if file.startswith("~$") or file.endswith(".tmp"):
                    continue  # Skip temporary office lock files
                full_file_path = os.path.join(root, file)
                if os.path.abspath(full_file_path) == abs_zip_path:
                    continue  # Skip zip file itself
                rel_path = os.path.relpath(full_file_path, base_output_dir)
                try:
                    zf.write(full_file_path, rel_path)
                except PermissionError:
                    pass

    print(f"Successfully created ZIP archive ({os.path.getsize(abs_zip_path)} bytes).")


import sys

sys.stdout.reconfigure(encoding="utf-8")

if __name__ == "__main__":
    results = run_batch_tests("outputs")

    # Generate Inputs Documentation in 3 formats
    generate_excel_inputs(TEST_CASES, "outputs/test_inputs.xlsx")
    generate_docx_inputs(TEST_CASES, "outputs/test_inputs.docx")
    generate_txt_inputs(TEST_CASES, "outputs/test_inputs.txt")

    # Bundle all outputs into a ZIP file
    create_outputs_zip("outputs", "outputs/all_test_outputs.zip")
    print("\nAll deliverables generated successfully!")

